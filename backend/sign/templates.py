"""
Template registry and PDF rendering for ChampPDF Sign.

Templates are legal instruments, not content. In v1 they are files in this
package (``templates/<id>.json`` + ``templates/<id>.body.html``), versioned
by git, loaded read-only at startup. A template's SHA-256 is recorded on every
document executed against it, so a later edit can never silently change what
a past document meant. The template-builder UI for Legal is deferred.

Two rendering paths:

  HTML body  (default)  ``<id>.body.html`` with ``{{merge_fields}}`` is laid
                        out by PyMuPDF's Story engine into an A4 PDF.
  DOCX master (preferred once Legal drops it in)
                        if ``<id>.docx`` exists next to the spec, its
                        ``{{placeholders}}`` are filled with python-docx and
                        the result is converted with LibreOffice. This is how
                        the real Master Template v1 .docx becomes the
                        rendered instrument without retyping it.

Both paths end with the same drawn execution page, which is where the
signature anchors live. The renderer returns those anchors so the provider
knows exactly where to stamp each party's signature.

First Schedule of the IT Act 2000: a template whose ``instrument_class`` is
one of the excluded classes is refused at load time. It cannot be sent by
accident because it cannot exist in the registry.
"""

from __future__ import annotations

import asyncio
import hashlib
import html
import io
import json
import logging
import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

logger = logging.getLogger(__name__)

TEMPLATE_DIR = Path(__file__).parent / "templates"

# IT Act 2000, First Schedule: documents to which the Act (and therefore
# electronic execution under s.3A / s.10A) does not apply.
FIRST_SCHEDULE_INSTRUMENTS = frozenset(
    {
        "negotiable_instrument",
        "power_of_attorney",
        "trust_deed",
        "will",
        "testamentary_disposition",
        "immovable_property_conveyance",
    }
)

PLACEHOLDER_RE = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}")

# Entity block. Defaults carry the NDA playbook's [CONFIRM] markers so an
# unconfirmed CIN can never be mistaken for a real one.
ENTITY_ENV = {
    "champions_entity": ("SIGN_ENTITY_NAME", "Champions Superior Capital"),
    "champions_cin": ("SIGN_ENTITY_CIN", "[CONFIRM: CIN]"),
    "champions_address": ("SIGN_ENTITY_ADDRESS", "[CONFIRM: registered address], Bengaluru, Karnataka, India"),
    "champions_signatory": ("SIGN_ENTITY_SIGNATORY", "Authorised Signatory"),
    "champions_signatory_designation": ("SIGN_ENTITY_SIGNATORY_DESIGNATION", "Director"),
}


class TemplateError(Exception):
    def __init__(self, message: str, field_key: Optional[str] = None) -> None:
        super().__init__(message)
        self.field_key = field_key


@dataclass
class TemplateField:
    key: str
    label: str
    type: str = "text"  # text | textarea | email | date | select
    required: bool = True
    placeholder: str = ""
    help: str = ""
    default: str = ""
    options: Optional[List[str]] = None
    max_length: int = 2000
    empty_text: str = ""  # rendered in place of an empty optional value (e.g. a [CONFIRM] marker)

    def to_public(self) -> Dict[str, Any]:
        d = {
            "key": self.key,
            "label": self.label,
            "type": self.type,
            "required": self.required,
            "placeholder": self.placeholder,
            "help": self.help,
            "default": date.today().isoformat() if self.default == "today" else self.default,
        }
        if self.options:
            d["options"] = list(self.options)
        return d


@dataclass
class TemplateRole:
    role: str  # signer | countersigner
    party: str  # counterparty | champions
    label: str


@dataclass
class Template:
    id: str
    version: int
    name: str
    short_name: str
    category: str
    instrument_class: str
    description: str
    title_format: str
    fields: List[TemplateField]
    roles: List[TemplateRole]
    body_html: str
    css: str
    schedule_defaults: Dict[str, str]
    docx_path: Optional[Path]
    sha256: str
    extra: Dict[str, Any] = field(default_factory=dict)

    @property
    def uses_docx(self) -> bool:
        return self.docx_path is not None and self.docx_path.exists()

    def to_public(self) -> Dict[str, Any]:
        return {
            "id": self.id,
            "version": self.version,
            "name": self.name,
            "short_name": self.short_name,
            "category": self.category,
            "instrument_class": self.instrument_class,
            "description": self.description,
            "fields": [f.to_public() for f in self.fields],
            "roles": [r.__dict__ for r in self.roles],
            "schedule_defaults": self.schedule_defaults,
            "source": "docx" if self.uses_docx else "html",
            "sha256": self.sha256,
        }

    def title_for(self, merge: Dict[str, str]) -> str:
        return PLACEHOLDER_RE.sub(lambda m: str(merge.get(m.group(1), "")), self.title_format).strip()


# --------------------------------------------------------------------------
# Registry
# --------------------------------------------------------------------------

_REGISTRY: Optional[Dict[str, Template]] = None


def _load_template(spec_path: Path) -> Template:
    raw = spec_path.read_bytes()
    spec = json.loads(raw.decode("utf-8"))
    instrument_class = spec.get("instrument_class", "")
    if instrument_class in FIRST_SCHEDULE_INSTRUMENTS:
        raise TemplateError(
            f"template {spec.get('id')} is a First Schedule instrument "
            f"({instrument_class}) and cannot be executed electronically"
        )
    body_path = spec_path.parent / spec.get("body", f"{spec['id']}.body.html")
    body_bytes = body_path.read_bytes() if body_path.exists() else b""
    css_path = spec_path.parent / spec.get("css", "base.css")
    css = css_path.read_text(encoding="utf-8") if css_path.exists() else ""
    docx_path = spec_path.parent / spec.get("docx", f"{spec['id']}.docx")
    docx_bytes = docx_path.read_bytes() if docx_path.exists() else b""

    h = hashlib.sha256()
    for chunk in (raw, body_bytes, css.encode("utf-8"), docx_bytes):
        h.update(hashlib.sha256(chunk).digest())

    fields = [
        TemplateField(
            key=f["key"],
            label=f["label"],
            type=f.get("type", "text"),
            required=bool(f.get("required", True)),
            placeholder=f.get("placeholder", ""),
            help=f.get("help", ""),
            default=str(f.get("default", "")),
            options=f.get("options"),
            max_length=int(f.get("max_length", 2000)),
            empty_text=str(f.get("empty_text", "")),
        )
        for f in spec.get("fields", [])
    ]
    roles = [TemplateRole(**r) for r in spec.get("roles", [])]
    if not any(r.role == "signer" for r in roles):
        raise TemplateError(f"template {spec.get('id')} declares no signer role")

    return Template(
        id=spec["id"],
        version=int(spec.get("version", 1)),
        name=spec["name"],
        short_name=spec.get("short_name", spec["name"]),
        category=spec.get("category", "agreement"),
        instrument_class=instrument_class,
        description=spec.get("description", ""),
        title_format=spec.get("title_format", spec["name"]),
        fields=fields,
        roles=roles,
        body_html=body_bytes.decode("utf-8"),
        css=css,
        schedule_defaults=spec.get("schedule_defaults", {}),
        docx_path=docx_path if docx_path.exists() else None,
        sha256=h.hexdigest(),
        extra={k: v for k, v in spec.items() if k not in {"fields", "roles"}},
    )


def load_registry(force: bool = False) -> Dict[str, Template]:
    global _REGISTRY
    if _REGISTRY is not None and not force:
        return _REGISTRY
    reg: Dict[str, Template] = {}
    if TEMPLATE_DIR.exists():
        for spec_path in sorted(TEMPLATE_DIR.glob("*.json")):
            try:
                t = _load_template(spec_path)
            except TemplateError as e:
                logger.error("Sign template %s refused: %s", spec_path.name, e)
                continue
            except Exception as e:  # noqa: BLE001 — a broken spec must not take the app down
                logger.error("Sign template %s failed to load: %s", spec_path.name, e)
                continue
            reg[t.id] = t
            logger.info("Sign template loaded: %s v%s (%s)", t.id, t.version, "docx" if t.uses_docx else "html")
    _REGISTRY = reg
    return reg


def get_template(template_id: str) -> Template:
    t = load_registry().get(template_id)
    if t is None:
        raise TemplateError(f"unknown template: {template_id}", "template_id")
    return t


def list_templates() -> List[Dict[str, Any]]:
    return [t.to_public() for t in load_registry().values()]


# --------------------------------------------------------------------------
# Merge fields
# --------------------------------------------------------------------------


def entity_fields() -> Dict[str, str]:
    return {key: os.environ.get(env, "").strip() or default for key, (env, default) in ENTITY_ENV.items()}


def validate_merge_fields(template: Template, values: Dict[str, Any]) -> Dict[str, str]:
    """Clean and validate sender-supplied values against the template's fields."""
    out: Dict[str, str] = {}
    for f in template.fields:
        raw = values.get(f.key)
        val = "" if raw is None else str(raw).strip()
        if not val and f.default:
            val = date.today().isoformat() if f.default == "today" else f.default
        if f.required and not val:
            raise TemplateError(f"{f.label} is required", f.key)
        if len(val) > f.max_length:
            raise TemplateError(f"{f.label} is too long (max {f.max_length} characters)", f.key)
        if f.type == "date" and val:
            try:
                val = date.fromisoformat(val).isoformat()
            except ValueError:
                raise TemplateError(f"{f.label} must be a date (YYYY-MM-DD)", f.key)
        if f.type == "select" and f.options and val and val not in f.options:
            raise TemplateError(f"{f.label} must be one of: {', '.join(f.options)}", f.key)
        if f.type == "email" and val and "@" not in val:
            raise TemplateError(f"{f.label} must be an email address", f.key)
        out[f.key] = val
    return out


def _display_date(value: str) -> str:
    try:
        return date.fromisoformat(value).strftime("%d %B %Y")
    except ValueError:
        return value


def build_merge_context(template: Template, values: Dict[str, str], recipients: List[Dict[str, Any]]) -> Dict[str, str]:
    """Everything the body may reference: validated fields, entity block, schedule, recipients."""
    ctx: Dict[str, str] = {}
    ctx.update(entity_fields())
    ctx.update({f"schedule_{k}": str(v) for k, v in template.schedule_defaults.items()})
    ctx.update(values)
    for f in template.fields:
        if not ctx.get(f.key) and f.empty_text:
            ctx[f.key] = f.empty_text
    for f in template.fields:
        if f.type == "date" and values.get(f.key):
            ctx[f"{f.key}_long"] = _display_date(values[f.key])
    for r in recipients:
        role = r.get("role", "signer")
        ctx[f"{role}_name"] = r.get("name", "")
        ctx[f"{role}_email"] = r.get("email", "")
        ctx[f"{role}_designation"] = r.get("designation", "") or ""
    ctx["template_id"] = template.id
    ctx["template_version"] = str(template.version)
    ctx["rendered_on"] = datetime.utcnow().strftime("%d %B %Y")
    return ctx


def substitute(text: str, ctx: Dict[str, str], escape: bool = True) -> str:
    def repl(m: "re.Match[str]") -> str:
        v = str(ctx.get(m.group(1), ""))
        return html.escape(v) if escape else v

    return PLACEHOLDER_RE.sub(repl, text)


# --------------------------------------------------------------------------
# Rendering
# --------------------------------------------------------------------------


@dataclass
class RenderResult:
    pdf: bytes
    anchors: Dict[str, Dict[str, Any]]  # role -> {page, rect, name_pos, designation_pos, date_pos}
    page_count: int
    source: str  # html | docx


A4 = (595.0, 842.0)
MARGIN = 54.0


def _story_to_pdf(html_doc: str, css: str) -> bytes:
    import pymupdf

    story = pymupdf.Story(html=html_doc, user_css=css or None)
    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    mediabox = pymupdf.Rect(0, 0, *A4)
    where = pymupdf.Rect(MARGIN, MARGIN, A4[0] - MARGIN, A4[1] - MARGIN - 18)
    more = 1
    guard = 0
    while more and guard < 200:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
        guard += 1
    writer.close()
    return buf.getvalue()


def _fill_docx(docx_bytes: bytes, ctx: Dict[str, str]) -> bytes:
    """Replace {{placeholders}} in paragraphs and table cells of a .docx."""
    import docx  # python-docx

    document = docx.Document(io.BytesIO(docx_bytes))

    def fix_paragraph(p) -> None:
        full = "".join(run.text for run in p.runs)
        if "{{" not in full:
            return
        new = substitute(full, ctx, escape=False)
        if new == full:
            return
        # Word splits text across runs unpredictably; collapse into the first
        # run (keeps its formatting) and blank the rest.
        if p.runs:
            p.runs[0].text = new
            for run in p.runs[1:]:
                run.text = ""

    for p in document.paragraphs:
        fix_paragraph(p)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    fix_paragraph(p)
    for section in document.sections:
        for part in (section.header, section.footer):
            for p in part.paragraphs:
                fix_paragraph(p)

    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _draw_execution_page(pdf_bytes: bytes, template: Template, ctx: Dict[str, str],
                         recipients: List[Dict[str, Any]], footer_label: str) -> RenderResult:
    """
    Append the execution page (signature blocks) and page footers. Returns the
    PDF plus the anchor rectangles for each role, in PyMuPDF coordinates
    (origin top-left, points).
    """
    import pymupdf

    doc = pymupdf.open("pdf", pdf_bytes)
    page = doc.new_page(width=A4[0], height=A4[1])
    y = MARGIN + 10
    page.insert_text((MARGIN, y), "EXECUTION", fontname="hebo", fontsize=13, color=(0.08, 0.1, 0.15))
    y += 22
    intro = (
        "IN WITNESS WHEREOF the Parties have executed this Agreement electronically through ChampPDF Sign "
        "on the dates recorded against each signature. Each signature is bound to the audit record "
        "referenced in the Certificate of Completion appended at execution, which records the identity "
        "verification, the exact document viewed, and the time and network address of signing."
    )
    rect = pymupdf.Rect(MARGIN, y, A4[0] - MARGIN, y + 70)
    page.insert_textbox(rect, intro, fontname="helv", fontsize=9.5, color=(0.25, 0.27, 0.33), lineheight=1.35)
    y += 84

    by_role = {r.get("role", "signer"): r for r in recipients}
    col_w = (A4[0] - 2 * MARGIN - 24) / 2
    anchors: Dict[str, Dict[str, Any]] = {}
    # Counterparty (signer) on the left, Champions (countersigner) on the right.
    ordered = sorted(template.roles, key=lambda r: 0 if r.party == "counterparty" else 1)
    for i, role in enumerate(ordered):
        x = MARGIN + i * (col_w + 24)
        entity = ctx.get("counterparty_entity", "") if role.party == "counterparty" else ctx.get("champions_entity", "")
        rec = by_role.get(role.role, {})
        page.insert_text((x, y), role.label.upper(), fontname="helv", fontsize=7.5, color=(0.4, 0.42, 0.5))
        page.insert_textbox(
            pymupdf.Rect(x, y + 6, x + col_w, y + 40),
            f"Signed for and on behalf of {entity}",
            fontname="hebo", fontsize=9.5, color=(0.08, 0.1, 0.15), lineheight=1.3,
        )
        box = pymupdf.Rect(x, y + 46, x + col_w, y + 46 + 74)
        page.draw_rect(box, color=(0.72, 0.75, 0.82), width=0.8, dashes="[3 2] 0")
        page.insert_text((x + 6, y + 46 + 12), "Signature", fontname="helv", fontsize=7, color=(0.55, 0.58, 0.65))
        ly = box.y1 + 20
        name = rec.get("name") or (ctx.get("champions_signatory", "") if role.party == "champions" else "")
        desig = rec.get("designation") or (ctx.get("champions_signatory_designation", "") if role.party == "champions" else "")
        lines = [("Name", name), ("Designation", desig), ("Date", "")]
        positions: Dict[str, List[float]] = {}
        for label, value in lines:
            page.insert_text((x, ly), f"{label}:", fontname="helv", fontsize=9, color=(0.4, 0.42, 0.5))
            vx = x + 68
            page.draw_line((vx, ly + 2), (x + col_w, ly + 2), color=(0.8, 0.82, 0.87), width=0.6)
            if value:
                page.insert_text((vx + 2, ly), value[:60], fontname="helv", fontsize=9.5, color=(0.08, 0.1, 0.15))
            positions[label.lower()] = [vx + 2, ly]
            ly += 20
        anchors[role.role] = {
            "page": page.number,
            "rect": [box.x0, box.y0, box.x1, box.y1],
            "name_pos": positions["name"],
            "designation_pos": positions["designation"],
            "date_pos": positions["date"],
            "party": role.party,
            "prefilled_name": bool(name),
            "prefilled_designation": bool(desig),
        }

    total = doc.page_count
    for p in doc:
        p.insert_text(
            (MARGIN, A4[1] - 30),
            f"{footer_label}   |   {template.short_name} v{template.version}   |   Page {p.number + 1} of {total}",
            fontname="helv", fontsize=7.5, color=(0.5, 0.52, 0.6),
        )
    out = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    return RenderResult(pdf=out, anchors=anchors, page_count=total, source="html")


async def render_pdf(template: Template, values: Dict[str, str], recipients: List[Dict[str, Any]],
                     footer_label: str) -> RenderResult:
    """Render the filled instrument. Heavy work runs off the event loop."""
    ctx = build_merge_context(template, values, recipients)
    source = "html"
    if template.uses_docx:
        from doc_converter import office_to_pdf, office_to_pdf_available

        if office_to_pdf_available():
            filled = await asyncio.to_thread(_fill_docx, template.docx_path.read_bytes(), ctx)
            base = await office_to_pdf(filled, f"{template.id}.docx")
            source = "docx"
        else:
            logger.warning("Template %s has a .docx master but LibreOffice is unavailable; using HTML body", template.id)
            base = await asyncio.to_thread(_story_to_pdf, _html_document(template, ctx), template.css)
    else:
        base = await asyncio.to_thread(_story_to_pdf, _html_document(template, ctx), template.css)
    result = await asyncio.to_thread(_draw_execution_page, base, template, ctx, recipients, footer_label)
    result.source = source
    return result


def _html_document(template: Template, ctx: Dict[str, str]) -> str:
    body = substitute(template.body_html, ctx, escape=True)
    return f"<html><body>{body}</body></html>"
